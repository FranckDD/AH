"""Build docs/adoption/adoption-tracking.xlsx (no macros, no external links)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = Path(__file__).resolve().parent / "adoption-tracking.xlsx"
HEAD = PatternFill("solid", fgColor="0F5E57")
INPUT = PatternFill("solid", fgColor="FFFBE6")
CALC = PatternFill("solid", fgColor="EEF1F2")
thin = Side(style="thin", color="C9D1D3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 45


def setup(ws, freeze, last_col, last_row):
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.print_area = "A1:%s%d" % (L(last_col), last_row)


# Instructions
ws = wb.active
ws.title = "Instructions"
ws["A1"] = "Adoption tracking, Glostone-Care"
ws["A1"].font = Font(bold=True, size=14, color="0F5E57")
ws["A2"] = "1. Fill the yellow cells only. Each service fills one line per week in Weekly indicators. Log every problem in Issues log. Tick skills in Training attendance."
ws["A3"] = "2. Read the sheets at the weekly 15-minute review. Update Go-No-go at the end of the pilot and of each extension."
ws["A4"] = "3. Gray cells are formulas, do not type in them. Do not write patient names or patient data in this file."
ws["A6"] = "Start date of week 1"
ws["A6"].font = Font(bold=True)
ws["B6"].fill = INPUT
ws["B6"].border = BORDER
ws["B6"].number_format = "yyyy-mm-dd"
ws["C6"] = "Type the date, for example 2026-11-02"
for r in (2, 3, 4):
    ws[f"A{r}"].alignment = WRAP
    ws.merge_cells(f"A{r}:H{r}")
    ws.row_dimensions[r].height = 32
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 40
ws.page_setup.orientation = "landscape"

# Weekly indicators
ws = wb.create_sheet("Weekly indicators")
cols = ["Week no.", "Week start date", "Service", "Staff trained", "Staff using the system daily",
        "Records created in system", "Records still on paper only", "Sales or entries double-entered",
        "Incidents", "Average time of a standard task [minutes]", "Satisfaction (1 to 5)", "Notes",
        "Total records", "Percent in system", "Percent of trained staff using it daily"]
widths = [8, 14, 24, 10, 14, 13, 13, 14, 10, 16, 13, 36, 10, 11, 15]
header(ws, 1, cols, widths)
services = ["Reception and cash desk", "Nursing", "Medical (doctors)", "Laboratory", "Pharmacy and stock",
            "Production", "Toxicology", "Spiritual care", "Management and admin"]
WEEKS = 8
r = 2
for w in range(1, WEEKS + 1):
    for s in services:
        ws.cell(r, 1, w)
        ws.cell(r, 2, '=IF(Instructions!$B$6="","",Instructions!$B$6+7*(A%d-1))' % r).number_format = "yyyy-mm-dd"
        ws.cell(r, 3, s)
        ws.cell(r, 13, '=IF(AND(F%d="",G%d=""),"",N(F%d)+N(G%d))' % (r, r, r, r))
        ws.cell(r, 14, '=IF(N(M%d)=0,"",F%d/M%d)' % (r, r, r)).number_format = "0%"
        ws.cell(r, 15, '=IF(N(D%d)=0,"",E%d/D%d)' % (r, r, r)).number_format = "0%"
        for c in range(1, 16):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.alignment = WRAP
            if 4 <= c <= 12:
                cell.fill = INPUT
            elif c in (2, 13, 14, 15):
                cell.fill = CALC
        r += 1
last = r - 1
ws.cell(r, 3, "Total (sums and averages)").font = Font(bold=True)
for c in (4, 5, 6, 7, 8, 9):
    ws.cell(r, c, "=SUM(%s2:%s%d)" % (L(c), L(c), last)).font = Font(bold=True)
ws.cell(r, 10, '=IF(COUNT(J2:J%d)=0,"",AVERAGE(J2:J%d))' % (last, last)).number_format = "0.0"
ws.cell(r, 11, '=IF(COUNT(K2:K%d)=0,"",AVERAGE(K2:K%d))' % (last, last)).number_format = "0.0"
ws.cell(r, 13, "=SUM(M2:M%d)" % last)
ws.cell(r, 14, '=IF(M%d=0,"",F%d/M%d)' % (r, r, r)).number_format = "0%"
ws.cell(r, 15, '=IF(D%d=0,"",E%d/D%d)' % (r, r, r)).number_format = "0%"
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                    showErrorMessage=True, errorTitle="Satisfaction", error="Type a whole number from 1 to 5.")
ws.add_data_validation(dv)
dv.add("K2:K%d" % last)
dvn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                     showErrorMessage=True, error="Type a number, zero or more.")
ws.add_data_validation(dvn)
for col in "DEFGHIJ":
    dvn.add("%s2:%s%d" % (col, col, last))
setup(ws, "D2", 15, r)

# Issues log
ws = wb.create_sheet("Issues log")
cols = ["Date", "Service", "Reporter", "Description (no patient data)", "Severity", "Owner", "Status", "Date fixed"]
header(ws, 1, cols, [12, 22, 16, 60, 12, 16, 14, 12])
N = 200
for r in range(2, N + 2):
    for c in range(1, 9):
        cell = ws.cell(r, c)
        cell.border = BORDER
        cell.alignment = WRAP
        cell.fill = INPUT
    ws.cell(r, 1).number_format = "yyyy-mm-dd"
    ws.cell(r, 8).number_format = "yyyy-mm-dd"
dv1 = DataValidation(type="list", formula1='"Low,Medium,High,Critical"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"Open,In progress,Fixed,Closed"', allow_blank=True)
dv3 = DataValidation(type="list", formula1='"' + ",".join(services) + '"', allow_blank=True)
for d, col in ((dv1, "E"), (dv2, "G"), (dv3, "B")):
    ws.add_data_validation(d)
    d.add("%s2:%s%d" % (col, col, N + 1))
ws["J1"] = "Open issues"
ws["K1"] = '=COUNTIF(G2:G%d,"Open")+COUNTIF(G2:G%d,"In progress")' % (N + 1, N + 1)
ws["J2"] = "Open critical"
ws["K2"] = ('=COUNTIFS(E2:E%d,"Critical",G2:G%d,"Open")+COUNTIFS(E2:E%d,"Critical",G2:G%d,"In progress")'
            % (N + 1, N + 1, N + 1, N + 1))
ws.column_dimensions["J"].width = 16
for c in ("J1", "J2"):
    ws[c].font = Font(bold=True)
setup(ws, "A2", 8, N + 1)

# Training attendance
ws = wb.create_sheet("Training attendance")
cols = ["Name", "Role", "Session date", "Trained by", "Skill 1", "Skill 2", "Skill 3", "Skill 4", "Skill 5",
        "Skill 6", "Skill 7", "Skills shown (count)", "Comment"]
header(ws, 1, cols, [26, 22, 13, 18, 8, 8, 8, 8, 8, 8, 8, 14, 34])
N = 150
roles = ("Secretary and cash desk,Nurse,Doctor,Laboratory technician,Pharmacy and stock,Production,"
         "Toxicology team,Spiritual counsellor,Promoter,Admin")
dvr = DataValidation(type="list", formula1='"%s"' % roles, allow_blank=True)
dvx = DataValidation(type="list", formula1='"X"', allow_blank=True)
ws.add_data_validation(dvr)
ws.add_data_validation(dvx)
dvr.add("B2:B%d" % (N + 1))
dvx.add("E2:K%d" % (N + 1))
for r in range(2, N + 2):
    ws.cell(r, 12, '=IF(COUNTIF(E%d:K%d,"X")=0,"",COUNTIF(E%d:K%d,"X"))' % (r, r, r, r))
    ws.cell(r, 3).number_format = "yyyy-mm-dd"
    for c in range(1, 14):
        cell = ws.cell(r, c)
        cell.border = BORDER
        cell.fill = CALC if c == 12 else INPUT
        if 5 <= c <= 12:
            cell.alignment = Alignment(horizontal="center")
ws["O1"] = ("Skill numbers follow the checklist of the role (section 4 of the plan). "
            "Put X when the person showed the skill. Leave the unused skill columns empty.")
ws["O1"].alignment = WRAP
ws.column_dimensions["O"].width = 50
setup(ws, "B2", 13, N + 1)

# Go-No-go
ws = wb.create_sheet("Go-No-go")
header(ws, 1, ["No.", "Criterion", "Target", "Current value", "Answer", "Comment"], [6, 62, 14, 14, 14, 40])
crit = [
    "At least [X] percent of the service staff are trained",
    "At least [X] percent of the service staff use the system every day",
    "Records still on paper only are at most [X] percent of the weekly records",
    "Sales or entries entered twice are at most [X] per week",
    "No critical issue is open (no data loss, no wrong money total)",
    "Average time of a standard task is at most [X] minutes",
    "Average satisfaction is at least [X] out of 5",
    "The outage drill was done and the outage sheet is in place",
    "Cash closing done every working day for [X] days, no unexplained difference (reception)",
    "Support person and deputy champion confirmed for the next step",
]
dvg = DataValidation(type="list", formula1='"Yes,No,Not yet"', allow_blank=True)
ws.add_data_validation(dvg)
for i, t in enumerate(crit, 2):
    ws.cell(i, 1, i - 1)
    ws.cell(i, 2, t)
    for c in range(1, 7):
        cell = ws.cell(i, c)
        cell.border = BORDER
        cell.alignment = WRAP
        if c >= 3:
            cell.fill = INPUT
    dvg.add("E%d" % i)
end = len(crit) + 1
ws.cell(end + 2, 2, "Overall decision").font = Font(bold=True)
ws.cell(end + 2, 5, '=IF(COUNTIF(E2:E%d,"Yes")=%d,"GO",IF(COUNTIF(E2:E%d,"No")>0,"NO-GO","NOT YET"))'
        % (end, len(crit), end)).font = Font(bold=True)
ws.cell(end + 3, 2, "Rule: GO only if every line is Yes. One No gives NO-GO. Otherwise NOT YET.")
ws.cell(end + 4, 2, "Service or step being decided: [SERVICE]    Date: [DATE]    Decided by: [NAME]")
setup(ws, "A2", 6, end + 4)

wb.save(OUT)
print(OUT)
