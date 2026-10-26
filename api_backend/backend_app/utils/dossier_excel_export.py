# api_backend/backend_app/utils/dossier_excel_export.py
import io
from typing import Any, Dict
from openpyxl import Workbook


LAB_STATUS_FR = {"pending": "En attente", "partial": "Partiel", "analyzed": "Analysé", "completed": "Terminé"}


def build_dossier_excel(dossier: Dict[str, Any]) -> bytes:
    """
    Un onglet par domaine PRESENT dans le dossier (dict.get, jamais un
    acces direct par cle - dossier_toxico/historique_spirituel sont
    absents pour medecin/nurse, chantier perimetre medical 2026-09-22 -
    ce fichier ne doit jamais casser cette garantie ni la contourner).
    """
    wb = Workbook()
    wb.remove(wb.active)

    patient = dossier.get("patient") or {}
    ws_patient = wb.create_sheet("Patient")
    ws_patient.append(["Champ", "Valeur"])
    for key in ("code_patient", "first_name", "last_name", "birth_date", "gender", "contact_phone"):
        ws_patient.append([key, patient.get(key)])

    if dossier.get("historique_medical") is not None:
        ws = wb.create_sheet("Historique medical")
        ws.append(["Date", "Motif", "Diagnostic"])
        for r in dossier["historique_medical"]:
            ws.append([r.get("consultation_date") or r.get("created_at"), r.get("motif_code"), r.get("diagnosis")])

    if dossier.get("prescriptions") is not None:
        ws = wb.create_sheet("Prescriptions")
        ws.append(["Médicament", "Dosage", "Fréquence", "Début", "Fin"])
        for p in dossier["prescriptions"]:
            ws.append([p.get("medication"), p.get("dosage"), p.get("frequency"), p.get("start_date"), p.get("end_date")])

    if dossier.get("hospitalisations") is not None:
        ws = wb.create_sheet("Hospitalisations")
        ws.append(["Admission", "Sortie", "Durée", "Motif d'admission", "Type de sortie", "Note de sortie", "Admis par", "Sortie prononcée par"])
        for h in dossier["hospitalisations"]:
            ws.append([
                h.get("admitted_at"), h.get("discharged_at"), h.get("duration_label"), h.get("admission_reason"),
                h.get("disposition_label"), h.get("discharge_note"), h.get("admitted_by_name"), h.get("discharged_by_name"),
            ])

    if dossier.get("historique_labo") is not None:
        ws = wb.create_sheet("Labo")
        ws.append(["Date", "Examen", "Statut"])
        for r in dossier["historique_labo"]:
            ws.append([r.get("test_date"), r.get("examen_name"), LAB_STATUS_FR.get(r.get("status"), r.get("status"))])

    if dossier.get("dossier_toxico") is not None:
        ws = wb.create_sheet("Toxicologie")
        toxico = dossier["dossier_toxico"]
        ws.append(["Champ", "Valeur"])
        for key in ("substance", "current_phase", "admission_date"):
            ws.append([key, toxico.get(key)])

    if dossier.get("historique_spirituel") is not None:
        ws = wb.create_sheet("Spirituel")
        ws.append(["Date", "Type", "Intervenant"])
        for r in dossier["historique_spirituel"]:
            ws.append([r.get("consultation_date"), r.get("type_consultation"), r.get("created_by_name")])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
