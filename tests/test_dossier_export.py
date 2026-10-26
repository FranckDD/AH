# tests/test_dossier_export.py
import io
from datetime import date

from openpyxl import load_workbook
from pypdf import PdfReader

from api_backend.backend_app.routes.auth import auth_endpoints
from api_backend.backend_app.routes.patient_dossier import patient_dossier_endpoint
from models.medical_record import MedicalRecord
from models.toxico import ToxicoDossier
from tests.conftest import create_test_user, create_test_patient, auth_headers

TEST_PASSWORD = "Correct123!"


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_export_dossier_pdf_medecin_ne_contient_pas_toxico(db_session, api_client):
    """Chantier exports (2026-09-23) : la garantie de cloisonnement du
    chantier perimetre medical (2026-09-22) doit se propager au fichier
    exporte, pas seulement a l'ecran."""
    medecin = create_test_user(db_session, "export_dossier_medecin", "medecin", password=TEST_PASSWORD)
    patient_id, _ = create_test_patient(db_session, medecin, first_name="ExportDossierMulti")
    db_session.flush()

    db_session.add(MedicalRecord(patient_id=patient_id, motif_code="free", diagnosis="RAS"))
    db_session.add(ToxicoDossier(patient_id=patient_id, admission_date=date(2026, 1, 10), substance="Alcool"))
    db_session.flush()

    client = api_client(auth_endpoints, patient_dossier_endpoint)
    headers = auth_headers(client, "export_dossier_medecin", TEST_PASSWORD)

    resp_pdf = client.get(f"/patients/{patient_id}/dossier/export?format=pdf", headers=headers)
    resp_excel = client.get(f"/patients/{patient_id}/dossier/export?format=excel", headers=headers)

    assert resp_pdf.status_code == 200
    assert resp_pdf.headers["content-type"] == "application/pdf"
    assert "Toxicologie" not in _pdf_text(resp_pdf.content)

    assert resp_excel.status_code == 200
    assert resp_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(io.BytesIO(resp_excel.content))
    assert "Toxicologie" not in wb.sheetnames
    assert "Spirituel" not in wb.sheetnames


def test_export_dossier_admin_contient_toxico(db_session, api_client):
    """Non-regression : seul medecin/nurse est concerne par le
    cloisonnement, admin garde tout."""
    admin = create_test_user(db_session, "export_dossier_admin", "admin", password=TEST_PASSWORD)
    patient_id, _ = create_test_patient(db_session, admin, first_name="ExportDossierAdminMulti")
    db_session.flush()

    db_session.add(MedicalRecord(patient_id=patient_id, motif_code="free", diagnosis="RAS"))
    db_session.add(ToxicoDossier(patient_id=patient_id, admission_date=date(2026, 1, 10), substance="Alcool"))
    db_session.flush()

    client = api_client(auth_endpoints, patient_dossier_endpoint)
    headers = auth_headers(client, "export_dossier_admin", TEST_PASSWORD)

    resp_pdf = client.get(f"/patients/{patient_id}/dossier/export?format=pdf", headers=headers)

    assert resp_pdf.status_code == 200
    assert "Toxicologie" in _pdf_text(resp_pdf.content)

    resp_excel = client.get(f"/patients/{patient_id}/dossier/export?format=excel", headers=headers)
    assert resp_excel.status_code == 200
    wb = load_workbook(io.BytesIO(resp_excel.content))
    assert "Toxicologie" in wb.sheetnames


def _stay(db_session, patient_id, user, reason, note, disposition="GUERI"):
    from datetime import datetime
    from models.hospitalization import Hospitalization
    db_session.add(Hospitalization(
        patient_id=patient_id, admitted_by=user.user_id, admitted_at=datetime(2026, 10, 1, 8, 30),
        admission_reason=reason, discharged_at=datetime(2026, 10, 4, 10, 0),
        discharge_disposition=disposition, discharge_note=note, discharged_by=user.user_id,
    ))
    db_session.flush()


def test_export_dossier_contient_les_sejours_et_notes_de_sortie(db_session, api_client):
    """Defaut releve par les guides 2026-10-26 : l'export du dossier (PDF et
    Excel) omettait les sejours hospitaliers et les notes de sortie."""
    medecin = create_test_user(db_session, "export_dossier_hosp", "medecin", password=TEST_PASSWORD)
    patient_id, _ = create_test_patient(db_session, medecin, first_name="ExportHosp")
    db_session.flush()
    db_session.add(MedicalRecord(patient_id=patient_id, motif_code="free", diagnosis="RAS"))
    _stay(db_session, patient_id, medecin, "Paludisme grave", "Controle dans une semaine")

    client = api_client(auth_endpoints, patient_dossier_endpoint)
    headers = auth_headers(client, "export_dossier_hosp", TEST_PASSWORD)

    pdf = client.get(f"/patients/{patient_id}/dossier/export?format=pdf", headers=headers)
    texte = " ".join(_pdf_text(pdf.content).split())
    assert "Hospitalisations" in texte
    assert "Paludisme grave" in texte
    assert "Controle dans une semaine" in texte
    assert "Guéri" in texte

    excel = client.get(f"/patients/{patient_id}/dossier/export?format=excel", headers=headers)
    wb = load_workbook(io.BytesIO(excel.content))
    assert "Hospitalisations" in wb.sheetnames
    lignes = [[c for c in row] for row in wb["Hospitalisations"].iter_rows(values_only=True)]
    assert any("Paludisme grave" in (c or "") for row in lignes for c in row if isinstance(c, str))
    assert any("Controle dans une semaine" in (c or "") for row in lignes for c in row if isinstance(c, str))


def test_export_dossier_pdf_sans_anglais_ni_tiret_long(db_session, api_client):
    medecin = create_test_user(db_session, "export_dossier_fr", "medecin", password=TEST_PASSWORD)
    patient_id, _ = create_test_patient(db_session, medecin, first_name="ExportFr")
    db_session.flush()
    db_session.add(MedicalRecord(patient_id=patient_id, motif_code="free", diagnosis="RAS"))
    db_session.flush()

    client = api_client(auth_endpoints, patient_dossier_endpoint)
    headers = auth_headers(client, "export_dossier_fr", TEST_PASSWORD)
    texte = _pdf_text(client.get(f"/patients/{patient_id}/dossier/export?format=pdf", headers=headers).content)

    assert "\u2014" not in texte
    assert "Page 1 sur" in texte
